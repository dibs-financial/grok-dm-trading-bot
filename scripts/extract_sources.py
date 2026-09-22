#!/usr/bin/env python3
"""Ingest DM Market Update memos/slides (PDF or PPTX) into knowledge/sources.

Usage:
    python3 scripts/extract_sources.py <file-or-dir> [...]

For every input file:
  * duplicates (identical MD5) are skipped,
  * the date and kind (memo|slides) are inferred from the filename,
  * the original is copied to knowledge/sources/original/<date>_<kind>.<ext>,
  * a plain-text rendering is written to knowledge/sources/text/<date>_<kind>.txt.

Requires: pymupdf (PDF), python-pptx (PPTX).
"""
import glob, hashlib, os, re, shutil, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIG = os.path.join(ROOT, "knowledge", "sources", "original")
TEXT = os.path.join(ROOT, "knowledge", "sources", "text")

MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}


def infer_date(name):
    """Return YYYY-MM-DD from names like 'memo_may11_2026', 'Slides_Aug10_2026', 'june21'."""
    m = re.search(r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*_?(\d{1,2})(?:_?(\d{4}))?",
                  name, re.I)
    if not m:
        return None
    month = MONTHS[m.group(1).lower()[:3]]
    day = int(m.group(2))
    year = int(m.group(3)) if m.group(3) else 2026
    return f"{year:04d}-{month:02d}-{day:02d}"


def infer_kind(name):
    return "memo" if "memo" in name.lower() else "slides"


def pdf_text(path):
    import pymupdf
    out = []
    for i, page in enumerate(pymupdf.open(path)):
        out.append(f"\n===== PAGE {i + 1} =====\n" + page.get_text())
    return "\n".join(out)


def pptx_text(path):
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(path).slides):
        out.append(f"\n===== SLIDE {i + 1} =====")
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
            if sh.shape_type == 13:
                out.append(f"[IMAGE {sh.name}]")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("--- NOTES ---\n" + slide.notes_slide.notes_text_frame.text)
    return "\n".join(out)


def main(args):
    os.makedirs(ORIG, exist_ok=True)
    os.makedirs(TEXT, exist_ok=True)
    files = []
    for a in args:
        files += sorted(glob.glob(os.path.join(a, "*"))) if os.path.isdir(a) else [a]
    # hashes already present in the repo
    seen = {hashlib.md5(open(p, "rb").read()).hexdigest(): os.path.basename(p)
            for p in glob.glob(os.path.join(ORIG, "*"))}
    for f in files:
        ext = f.rsplit(".", 1)[-1].lower()
        if ext not in ("pdf", "pptx"):
            continue
        h = hashlib.md5(open(f, "rb").read()).hexdigest()
        if h in seen:
            print(f"DUP   {os.path.basename(f)} == {seen[h]}")
            continue
        base = os.path.basename(f)
        base = base.split("-", 1)[1] if re.match(r"^[0-9a-f]{8}-", base) else base
        date, kind = infer_date(base), infer_kind(base)
        if not date:
            print(f"SKIP  {base}: cannot infer date"); continue
        stem = f"{date}_{kind}"
        n = 2
        while os.path.exists(os.path.join(ORIG, f"{stem}.{ext}")) or \
                os.path.exists(os.path.join(TEXT, f"{stem}.txt")):
            stem = f"{date}_{kind}_v{n}"; n += 1
        shutil.copy(f, os.path.join(ORIG, f"{stem}.{ext}"))
        os.chmod(os.path.join(ORIG, f"{stem}.{ext}"), 0o644)
        text = pdf_text(f) if ext == "pdf" else pptx_text(f)
        with open(os.path.join(TEXT, f"{stem}.txt"), "w") as fh:
            fh.write(text)
        seen[h] = f"{stem}.{ext}"
        print(f"OK    {base} -> {stem} ({len(text)} chars)")


if __name__ == "__main__":
    main(sys.argv[1:] or ["."])
